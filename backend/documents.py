"""Private file ingestion, cited extraction and immutable PDF/ZIP exports."""
import hashlib
import io
import json
import os
import shutil
import subprocess
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from pypdf import PdfReader, PdfWriter
from docx import Document
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, PageBreak, Image as PDFImage
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from .domain import DomainError
from .malware import malware_screen

MAX_FILE = 10 * 1024 * 1024
MAX_TEXT = 200_000


def extract_file(name, content):
    if not content or len(content)>MAX_FILE: raise DomainError('Upload a non-empty file up to 10 MiB')
    scan_status=malware_screen(content)
    name=Path(name.replace('\\','/')).name[:180]
    if not name: raise DomainError('A file name is required')
    ext=Path(name).suffix.lower();text='';status='extracted'
    try:
        if ext=='.pdf' and content.startswith(b'%PDF-'):
            mime='application/pdf';reader=PdfReader(io.BytesIO(content))
            if reader.is_encrypted: raise DomainError('Upload an unencrypted PDF')
            if len(reader.pages)>100: raise DomainError('Split PDFs into files of at most 100 pages')
            chunks=[]
            for n,page in enumerate(reader.pages,1):
                # Bound extraction output and page count before it enters the database.
                raw=(page.extract_text() or '')[:MAX_TEXT]
                chunks.extend(f'[page {n}, line {j}] {line}' for j,line in enumerate(raw.splitlines(),1))
                if sum(map(len,chunks))>MAX_TEXT:break
            text='\n'.join(chunks)[:MAX_TEXT]
            if not text.strip():status='needs-ocr'
        elif ext=='.docx' and content.startswith(b'PK'):
            mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist())>30*1024*1024: raise DomainError('Expanded DOCX is too large')
            doc=Document(io.BytesIO(content));chunks=[f'[paragraph {n}] {p.text}' for n,p in enumerate(doc.paragraphs,1) if p.text.strip()]
            for n,table in enumerate(doc.tables,1):
                chunks.extend(f'[table {n}, row {j}] '+' | '.join(c.text for c in row.cells) for j,row in enumerate(table.rows,1))
            text='\n'.join(chunks)[:MAX_TEXT]
        elif ext in ('.png','.jpg','.jpeg'):
            im=Image.open(io.BytesIO(content));im.verify()
            if im.width*im.height>25_000_000:raise DomainError('Images must be at most 25 megapixels')
            if im.format not in ('PNG','JPEG'):raise DomainError('Image content does not match a supported format')
            mime='image/png' if im.format=='PNG' else 'image/jpeg';status='needs-ocr'
        elif ext in ('.txt','.csv'):
            mime='text/plain';text=content.decode('utf-8-sig')
            if '\x00' in text:raise DomainError('Binary content is not a text file')
            text='\n'.join(f'[line {n}] {line}' for n,line in enumerate(text.splitlines(),1))[:MAX_TEXT]
        else: raise DomainError('Supported files: PDF, DOCX, PNG, JPG, UTF-8 TXT and CSV')
    except DomainError:raise
    except Exception as exc:raise DomainError('File could not be read; upload a valid, unencrypted document') from exc
    return {'name':name,'mime':mime,'size':len(content),'checksum':hashlib.sha256(content).hexdigest(),'extracted':text,'extraction_status':status,'scan_status':scan_status}


def register_file_controls(db,file_id,scan_status):
    from datetime import datetime,timedelta,timezone
    days=int(os.environ.get('DOAIDE_FILE_RETENTION_DAYS','0'))
    expiry=(datetime.now(timezone.utc)+timedelta(days=days)).isoformat() if days>0 else None
    db.execute('INSERT OR REPLACE INTO file_controls(file_id,scan_status,retain_until) VALUES (?,?,?)',(file_id,scan_status,expiry))


def ocr_file(file):
    if not shutil.which('tesseract'):raise DomainError('Install Tesseract locally to enable OCR')
    if file['mime']=='application/pdf' and not shutil.which('pdftoppm'):raise DomainError('Install Poppler locally to OCR PDFs')
    with tempfile.TemporaryDirectory(prefix='doaide-ocr-') as directory:
        root=Path(directory);source=root/('input.pdf' if file['mime']=='application/pdf' else 'input.png');source.write_bytes(file['content'])
        if file['mime']=='application/pdf':
            reader=PdfReader(io.BytesIO(file['content']))
            if len(reader.pages)>20:raise DomainError('OCR accepts up to 20 pages per file')
            subprocess.run(['pdftoppm','-scale-to','1800','-png',str(source),str(root/'page')],capture_output=True,timeout=90,check=True)
            pages=sorted(root.glob('page-*.png'))
        elif file['mime'].startswith('image/'):pages=[source]
        else:raise DomainError('OCR is for PDF and image files')
        chunks=[]
        for n,page in enumerate(pages,1):
            output=subprocess.run(['tesseract',str(page),'stdout'],capture_output=True,text=True,timeout=30,check=True).stdout
            chunks.extend(f'[OCR page {n}, line {j}] {line}' for j,line in enumerate(output.splitlines(),1) if line.strip())
        return '\n'.join(chunks)[:MAX_TEXT]


def pdf_document(document, public=False, photos=None):
    snap=json.loads(document['snapshot']) if isinstance(document['snapshot'],str) else document['snapshot']
    output=io.BytesIO();styles=getSampleStyleSheet()
    # A Unicode font keeps customer names readable on macOS and Linux.
    candidates=['/System/Library/Fonts/Supplemental/Arial Unicode.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    font='Helvetica'
    for path in candidates:
        if Path(path).exists():
            if 'DoAide' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('DoAide',path))
            font='DoAide';break
    for style in styles.byName.values():style.fontName=font
    styles['BodyText'].fontSize=9;styles['BodyText'].leading=13;styles['BodyText'].spaceAfter=5
    styles['Heading2'].fontSize=12;styles['Heading2'].textColor=colors.HexColor('#365544')
    label=document['kind']=='traceability-label'
    if label:
        styles['BodyText'].fontSize=8;styles['BodyText'].leading=11
        styles['Heading2'].fontSize=10
    title=ParagraphStyle('DoAideTitle',parent=styles['Title'],fontSize=14 if label else 23,leading=18 if label else 29,alignment=TA_LEFT,textColor=colors.HexColor('#20382b'),spaceAfter=8 if label else 12)
    def p(value,style='BodyText'):return Paragraph(escape(str(value)).replace('\n','<br/>'),styles[style])
    story=[p('DOAIDE  /  '+document['kind'].replace('-',' ').upper(),'Heading2'),Paragraph(escape(snap['extra']['product'] if label else snap['record']['title']),title),p(document['number']+' · '+document['created_at'][:10]),p('Source version '+str(document['record_version'])+' · SHA-256 '+document['checksum']),Spacer(1,10)]
    def section(label,value):
        story.append(p(label,'Heading2'))
        render(value)
    def render(value,depth=0):
        if isinstance(value,dict):
            for key,val in value.items():
                if val in ('',None) or (key.startswith('_') and key not in ('_status','_waiver','_repair_note','_verification','_notes','_verified_at','_warranty_note')) or key in ('user_id','created_by','updated_by','record_id','answer','file_id'):continue
                label=key.replace('_',' ').capitalize()
                if isinstance(val,(dict,list)):
                    if val:story.append(p(label,'Heading3'));render(val,depth+1)
                else:
                    if key.endswith('_paise'):label=label.removesuffix(' paise');val=f'INR {val/100:,.2f}'
                    story.append(p(f'{label}: {val}'))
        elif isinstance(value,list):
            for n,item in enumerate(value,1):
                if isinstance(item,dict):
                    story.append(p(f'Item {n}','Heading3'));render(item,depth+1)
                else:story.append(p(item))
        else:story.append(p(value))
    if document['kind']=='traceability-label':
        section('Batch traceability',snap['extra'])
        story.append(p('Confirm current release / recall status in DoAide before dispatch. This label records status at issue time.'))
    elif document['kind']=='amc-renewal-offer':
        section('Customer',snap['payload']['customer'])
        offer=snap['extra']
        section('Proposed service period and terms',{'start':offer['start'],'end':offer['end'],'included_visits':offer['included'],'contract_value_paise':int(Decimal(str(offer['value']))*100),'approval_contact':offer['contact'],'offer_valid_until':offer['valid_until'],'terms':offer['terms']})
        story.append(p('Acceptance applies to this exact offer version. The existing contract continues under its own terms.'))
    elif document['kind']=='handover':
        section('Site',snap['payload'].get('site',snap['record']['title']))
        grouped={}
        for item in snap['items']:
            if item['kind']=='issue':grouped.setdefault(item['data']['room'],[]).append(item['data'])
        order={'Critical':0,'Major':1,'Minor':2}
        for room_number,(room,issues) in enumerate(sorted(grouped.items())):
            if room_number:story.append(PageBreak())
            story.append(p(room,'Heading2'))
            grid=[[p(x) for x in ('Issue','Severity','Status','Verification / waiver')]]
            for issue in sorted(issues,key=lambda x:order.get(x['severity'],3)):
                grid.append([p(issue['description']),p(issue['severity']),p(issue.get('_status','open')),p(issue.get('_verification') or issue.get('_waiver') or issue.get('_repair_note',''))])
            table=Table(grid,colWidths=[210,65,70,166],repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eee5')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#d6ded1'))]))
            story.extend([table,Spacer(1,10)])
            for issue in issues:
                for key in ('file_id','_repair_file'):
                    photo=(photos or {}).get(issue.get(key))
                    if not photo:continue
                    image=Image.open(io.BytesIO(photo));width,height=image.size;scale=min(245/width,180/height,1)
                    story.append(KeepTogether([p(issue['description']+' · '+('repair evidence' if key=='_repair_file' else 'inspection evidence')),PDFImage(io.BytesIO(photo),width=width*scale,height=height*scale),Spacer(1,10)]))
        story.append(p('Minor waivers listed here require client acceptance of this exact handover. Original evidence is retained in the authorized evidence packet.'))
    elif document['kind']=='fee-schedule-amendment':
        financial=snap['financial'];amendment=snap['extra'];payload=snap['payload']
        section('Agreement amendment',{'student':payload['student'],'course':payload['course'],'amendment_number':amendment['number'],'approved_by':amendment['approved_by'],'approved_at':amendment['at'][:10],'reason':amendment['reason']})
        def schedule_table(label,rows):
            story.append(p(label+' · INR','Heading2'))
            grid=[[p(v) for v in ('Due date','Amount','Allocated','Outstanding')]]
            for row in rows:
                grid.append([p(row['due'])]+[p(f"{row[k]/100:,.2f}") for k in ('amount_paise','allocated_paise','outstanding_paise')])
            table=Table(grid,colWidths=[145,120,120,120],repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eee5')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#d6ded1'))]))
            story.append(table);story.append(Spacer(1,10))
        schedule_table('Previous instalments',amendment['previous_allocations'])
        schedule_table('Revised instalments',financial['allocations'])
        section('Current balance', {k:financial[k] for k in ('total_paise','cash_received_paise','balance_paise','due_now_paise')})
        story.append(p('This amendment changes unpaid instalment dates and amounts. Previously issued receipts and approved financial entries remain unchanged.'))
    else:
        payload=dict(snap['payload'])
        if any(i['kind']=='estimate' for i in snap.get('items',[])):
            payload.pop('hours',None);payload.pop('rate',None)
        if any(i['kind']=='issue' for i in snap.get('items',[])):payload.pop('issues',None)
        if document['kind']=='receipt':
            payload={k:v for k,v in payload.items() if k in ('student','course','customer','invoice')}
            entry=next((e for e in snap.get('financial',{}).get('entries',[]) if e['id']==snap.get('extra',{}).get('entry_id')),None)
            if entry:
                section('Receipt / adjustment',{'reference':entry['reference'],'type':entry['kind'],'amount_paise':entry['amount'],'method':entry.get('method'),'date':entry['date'],'source':entry['note'],'recorded_by':entry['actor']})
        section('Record details',payload)
        if snap.get('financial'):section('Financial statement · INR',{k:v for k,v in snap['financial'].items() if k!='entries'})
        grouped={}
        for item in snap.get('items',[]):
            if public and item['kind'] in ('policy','guardian','assignment'):continue
            grouped.setdefault(item['kind'],[]).append(item['data'])
        for kind,rows in grouped.items():
            if document['kind']=='receipt':continue
            if kind=='estimate':
                story.append(p('Change estimate · INR','Heading2'))
                grid=[[p(v) for v in ('Work','Quantity','Unit rate','Discount','Tax %')]]
                grid.extend([[p(row[k]) for k in ('description','quantity','rate','discount','tax')] for row in rows])
                table=Table(grid,colWidths=[235,60,75,70,65],repeatRows=1,hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eee5')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#d6ded1'))]));story.append(table)
            else:section(kind.replace('_',' ').title(),rows)
        section('Workflow outcome',snap.get('workflow',{}))
        if snap.get('extra'):section('Document details',snap['extra'])
        if snap.get('files') and not public:section('Evidence manifest',snap['files'])
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont(font,6 if label else 8);canvas.setFillColor(colors.HexColor('#68756d'));canvas.drawString(18 if label else 42,12 if label else 28,document['number']);canvas.drawRightString(265 if label else 553,12 if label else 28,f'Page {doc.page}');canvas.restoreState()
    SimpleDocTemplate(output,pagesize=(283.46,425.20) if label else (595,842),rightMargin=18 if label else 42,leftMargin=18 if label else 42,topMargin=18 if label else 42,bottomMargin=26 if label else 48,title=snap['record']['title'],author='DoAide').build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()


def packet_zip(db,document):
    snap=json.loads(document['snapshot']);out=io.BytesIO();manifest=[]
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(document['number']+'.pdf',stored_pdf(db,document))
        for f in snap['files']:
            row=db.execute('SELECT * FROM files WHERE id=? AND record_id=?',(f['id'],document['record_id'])).fetchone()
            if not row:raise DomainError('A file in this immutable manifest is unavailable')
            if not row['content']:raise DomainError('Evidence was removed under the recorded retention policy; the immutable manifest remains available')
            archive.writestr('evidence/'+row['id'][:8]+'-'+row['name'],row['content']);manifest.append({k:row[k] for k in ('id','name','checksum','size')})
        archive.writestr('manifest.json',json.dumps({'document':document['number'],'snapshot_sha256':document['checksum'],'files':manifest,'approved_attachment_redactions':snap.get('extra',{}).get('approved_attachment_redactions',[])},indent=2))
    return out.getvalue()


def freeze_pdf(db,document):
    photos={}
    if document['kind']=='handover':
        for row in db.execute("SELECT id,content FROM files WHERE record_id=? AND mime LIKE 'image/%'",(document['record_id'],)):
            if row['content']:photos[row['id']]=row['content']
    content=pdf_document(document,photos=photos)
    snap=json.loads(document['snapshot'])
    if document['kind']=='claim-packet':
        writer=PdfWriter();writer.append(PdfReader(io.BytesIO(content)))
        for f in snap['files']:
            row=db.execute('SELECT * FROM files WHERE id=? AND record_id=?',(f['id'],document['record_id'])).fetchone()
            if not row or not row['content']:raise DomainError('A required evidence file is unavailable')
            if row['mime']=='application/pdf':
                writer.append(PdfReader(io.BytesIO(row['content'])))
            else:
                out=io.BytesIO();styles=getSampleStyleSheet();story=[Paragraph(escape(row['name']),styles['Heading1'])]
                if row['mime'].startswith('image/'):
                    source=Image.open(io.BytesIO(row['content']));width,height=source.size;scale=min(500/width,680/height,1)
                    story.append(PDFImage(io.BytesIO(row['content']),width=width*scale,height=height*scale))
                else:
                    story.append(Paragraph('Extracted evidence copy; the original is included in the evidence ZIP.',styles['BodyText']))
                    story.extend(Paragraph(escape(line),styles['BodyText']) for line in row['extracted'].splitlines())
                SimpleDocTemplate(out).build(story);writer.append(PdfReader(io.BytesIO(out.getvalue())))
        writer.add_metadata({'/Title':snap['record']['title'],'/Author':'DoAide'})
        out=io.BytesIO();writer.write(out);content=out.getvalue()
        template=next(i['data'] for i in snap['items'] if i['kind']=='template')
        if len(content)>int(template['max_mb'])*1024*1024:raise DomainError('Generated PDF exceeds the provider packet size limit')
    db.execute('INSERT INTO document_blobs VALUES (?,?,?)',(document['id'],content,hashlib.sha256(content).hexdigest()))
    public=pdf_document(document,public=True,photos=photos)
    db.execute('INSERT OR IGNORE INTO public_document_blobs VALUES (?,?,?)',(document['id'],public,hashlib.sha256(public).hexdigest()))
    return content


def stored_pdf(db,document):
    row=db.execute('SELECT pdf FROM document_blobs WHERE document_id=?',(document['id'],)).fetchone()
    return row['pdf'] if row else freeze_pdf(db,document)


def public_pdf(db,document):
    row=db.execute('SELECT pdf FROM public_document_blobs WHERE document_id=?',(document['id'],)).fetchone()
    if row:return row['pdf']
    content=pdf_document(document,public=True)
    db.execute('INSERT OR IGNORE INTO public_document_blobs VALUES (?,?,?)',(document['id'],content,hashlib.sha256(content).hexdigest()))
    return db.execute('SELECT pdf FROM public_document_blobs WHERE document_id=?',(document['id'],)).fetchone()['pdf']
